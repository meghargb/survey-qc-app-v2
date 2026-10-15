"""
Generates downloadable reports of QC results in the format the user picks.
"""
import io
import csv
import pandas as pd
from openpyxl import Workbook
from docx import Document
from pptx import Presentation
from pptx.util import Inches, Pt


def _flatten_for_table(results: dict) -> list[dict]:
    """Turns the nested results dict into flat rows for CSV/Excel/Word tables."""
    rows = []
    for c in results.get("completion", {}).get("details", []):
        rows.append({"Check": "Completion", "Item": c.get("question", ""),
                      "Detail": f"{c.get('blank_pct', '')}% blank",
                      "Flag": "YES" if c.get("flag") else "no"})
    for s in results.get("straight_lining", {}).get("details", []):
        rows.append({"Check": "Straight-lining", "Item": s.get("grid", ""),
                      "Detail": f"{s.get('pct', '')}% straight-liners",
                      "Flag": "YES" if s.get("flag") else "no"})
    sp = results.get("speeding", {})
    if "speeders_pct" in sp:
        rows.append({"Check": "Speeding", "Item": "All respondents",
                      "Detail": f"{sp.get('speeders_pct')}% speeding",
                      "Flag": "YES" if sp.get("flag") else "no"})
    for q in results.get("quota", {}).get("details", []):
        rows.append({"Check": "Quota", "Item": f"{q.get('quota_group')}={q.get('quota_label')}",
                      "Detail": f"{q.get('filled')}/{q.get('target_n')} ({q.get('fill_pct')}%)",
                      "Flag": q.get("status", "")})
    for w in results.get("wave_spike", {}).get("details", []):
        rows.append({"Check": "Wave Spike", "Item": w.get("question", ""),
                      "Detail": f"{w.get('previous_wave')} -> {w.get('current_wave')} ({w.get('difference'):+})"
                      if isinstance(w.get('difference'), (int, float)) else "",
                      "Flag": "YES" if w.get("flag") else "no"})
    for d in results.get("duplicates", {}).get("details", []):
        rows.append({"Check": "Duplicates", "Item": d.get("field", ""),
                      "Detail": f"{d.get('affected_respondents', '')} affected ({d.get('pct', '')}%)",
                      "Flag": "YES" if d.get("flag") else "no"})
    if results.get("duplicates", {}).get("full_row_duplicates"):
        rows.append({"Check": "Duplicates", "Item": "Full row duplicates",
                      "Detail": f"{results['duplicates']['full_row_duplicates']} rows "
                                f"({results['duplicates']['full_row_duplicate_pct']}%)",
                      "Flag": "YES" if results["duplicates"].get("flag") else "no"})
    for o in results.get("open_end_quality", {}).get("details", []):
        rows.append({"Check": "Open-End Quality", "Item": o.get("question", ""),
                      "Detail": f"{o.get('too_short_pct', '')}% too short, "
                                f"{o.get('copy_pasted_pct', '')}% copy-pasted",
                      "Flag": "YES" if o.get("flag") else "no"})
    for lc in results.get("logical_consistency", {}).get("details", []):
        rows.append({"Check": "Logical Consistency", "Item": lc.get("label", ""),
                      "Detail": f"{lc.get('contradictions', '')}/{lc.get('base', '')} "
                                f"({lc.get('pct', '')}%)",
                      "Flag": "YES" if lc.get("flag") else "no"})
    for ol in results.get("outliers", {}).get("details", []):
        rows.append({"Check": "Outliers", "Item": ol.get("column", ""),
                      "Detail": f"{ol.get('out_of_range_count', '')} out of range "
                                f"({ol.get('out_of_range_pct', '')}%), expected {ol.get('valid_range', '')}",
                      "Flag": "YES" if ol.get("flag") else "no"})
    for sl in results.get("skip_logic", {}).get("details", []):
        rows.append({"Check": "Skip Logic", "Item": sl.get("label", ""),
                      "Detail": f"{sl.get('violations', '')}/{sl.get('base', '')} "
                                f"({sl.get('pct', '')}%)",
                      "Flag": "YES" if sl.get("flag") else "no"})
    for rd in results.get("response_distribution", {}).get("details", []):
        rows.append({"Check": "Response Distribution", "Item": rd.get("question", ""),
                      "Detail": f"'{rd.get('top_value', '')}' at {rd.get('top_value_pct', '')}%",
                      "Flag": "YES" if rd.get("flag") else "no"})
    for at in results.get("attention_traps", {}).get("details", []):
        rows.append({"Check": "Attention Trap", "Item": at.get("label", ""),
                      "Detail": f"{at.get('failed', '')}/{at.get('base', '')} failed "
                                f"({at.get('pct', '')}%)",
                      "Flag": "YES" if at.get("flag") else "no"})
    return rows


def export_csv(study_name, results, summary, email_draft) -> bytes:
    buf = io.StringIO()
    writer = csv.DictWriter(buf, fieldnames=["Check", "Item", "Detail", "Flag"])
    writer.writeheader()
    writer.writerows(_flatten_for_table(results))
    return buf.getvalue().encode("utf-8")


def export_xlsx(study_name, results, summary, email_draft) -> bytes:
    wb = Workbook()
    ws = wb.active
    ws.title = "QC Results"
    ws.append(["Check", "Item", "Detail", "Flag"])
    for row in _flatten_for_table(results):
        ws.append([row["Check"], row["Item"], row["Detail"], row["Flag"]])

    ws2 = wb.create_sheet("Summary")
    ws2.append(["AI Summary"])
    ws2.append([summary])
    ws2.append([])
    ws2.append(["Vendor Email Draft"])
    ws2.append([email_draft])

    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def export_docx(study_name, results, summary, email_draft) -> bytes:
    doc = Document()
    doc.add_heading(f"Data QC Report — {study_name}", level=1)
    doc.add_paragraph(f"Checkpoint: {results.get('checkpoint_pct', '')}%")

    doc.add_heading("Summary", level=2)
    doc.add_paragraph(summary)

    doc.add_heading("Issues Found", level=2)
    issues = results.get("issues", [])
    if not issues:
        doc.add_paragraph("No issues detected.")
    else:
        for i in issues:
            doc.add_paragraph(i, style="List Bullet")

    doc.add_heading("Full Check Results", level=2)
    rows = _flatten_for_table(results)
    table = doc.add_table(rows=1, cols=4)
    table.style = "Light Grid Accent 1"
    hdr = table.rows[0].cells
    hdr[0].text, hdr[1].text, hdr[2].text, hdr[3].text = "Check", "Item", "Detail", "Flag"
    for r in rows:
        cells = table.add_row().cells
        cells[0].text, cells[1].text, cells[2].text, cells[3].text = (
            r["Check"], r["Item"], r["Detail"], r["Flag"])

    doc.add_heading("Draft Vendor Email", level=2)
    doc.add_paragraph(email_draft)

    buf = io.BytesIO()
    doc.save(buf)
    return buf.getvalue()


def export_pptx(study_name, results, summary, email_draft) -> bytes:
    prs = Presentation()

    slide = prs.slides.add_slide(prs.slide_layouts[0])
    slide.shapes.title.text = f"Data QC Report — {study_name}"
    slide.placeholders[1].text = f"Checkpoint: {results.get('checkpoint_pct', '')}%"

    slide2 = prs.slides.add_slide(prs.slide_layouts[1])
    slide2.shapes.title.text = "Summary"
    slide2.placeholders[1].text = summary

    slide3 = prs.slides.add_slide(prs.slide_layouts[1])
    slide3.shapes.title.text = "Issues Found"
    body = slide3.placeholders[1].text_frame
    issues = results.get("issues", [])
    if not issues:
        body.text = "No issues detected."
    else:
        body.text = issues[0]
        for i in issues[1:]:
            p = body.add_paragraph()
            p.text = i

    buf = io.BytesIO()
    prs.save(buf)
    return buf.getvalue()


def export_sav(study_name, df: pd.DataFrame) -> bytes:
    """Exports the underlying respondent-level data as an SPSS .sav file."""
    import pyreadstat
    import tempfile, os
    with tempfile.NamedTemporaryFile(suffix=".sav", delete=False) as tmp:
        tmp_path = tmp.name
    pyreadstat.write_sav(df, tmp_path)
    with open(tmp_path, "rb") as f:
        data = f.read()
    os.remove(tmp_path)
    return data
